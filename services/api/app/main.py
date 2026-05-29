from pathlib import Path
from typing import Any

from fastapi import FastAPI, HTTPException
from openpyxl import load_workbook
from pydantic import BaseModel, Field

app = FastAPI(
    title="FinLens Processing API",
    version="0.1.0",
    description="Deterministic parsing, validation, ratio, lineage, and AI-blueprint service."
)


class UploadSession(BaseModel):
    organisation_id: str
    file_name: str
    content_type: str
    document_type_hint: str | None = None


class UploadSessionResponse(BaseModel):
    upload_id: str
    signed_url: str
    expires_in_seconds: int = 900


class InsightBlueprint(BaseModel):
    workspace_id: str
    privacy_boundary: str = "masked-kpi-blueprint"
    masked_kpis: list[dict] = Field(default_factory=list)
    never_included: list[str] = Field(
        default_factory=lambda: [
            "raw rupee transaction rows",
            "customer or vendor legal names",
            "payroll identities",
            "bank account numbers",
            "uploaded file contents",
        ]
    )


class WorkbookInspectRequest(BaseModel):
    file_path: str
    max_preview_rows: int = Field(default=3, ge=0, le=10)


class SheetSummary(BaseModel):
    name: str
    rows: int
    columns: int
    headers: list[str]
    preview: list[dict[str, Any]]


class WorkbookInspectResponse(BaseModel):
    file_name: str
    sheets: list[SheetSummary]
    warnings: list[str] = Field(default_factory=list)


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok", "service": "finlens-api"}


@app.post("/uploads/sessions")
def create_upload_session(payload: UploadSession) -> UploadSessionResponse:
    return UploadSessionResponse(
        upload_id=f"upl_{payload.organisation_id}_{payload.file_name.replace('.', '_')}",
        signed_url="http://localhost:3000/api/mock-signed-upload-url",
    )


@app.post("/llm/blueprints/validate")
def validate_blueprint(payload: InsightBlueprint) -> dict[str, object]:
    forbidden_markers = ["invoice row", "customer name", "vendor name", "salary", "bank account"]
    serialized = payload.model_dump_json().lower()

    return {
        "valid": not any(marker in serialized for marker in forbidden_markers),
        "privacy_boundary": payload.privacy_boundary,
        "checked_markers": forbidden_markers,
    }


@app.post("/workbooks/inspect")
def inspect_workbook(payload: WorkbookInspectRequest) -> WorkbookInspectResponse:
    workbook_path = Path(payload.file_path).expanduser().resolve()
    if not workbook_path.exists() or not workbook_path.is_file():
        raise HTTPException(status_code=404, detail="Workbook file not found")

    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Unable to parse workbook") from exc

    warnings: list[str] = []
    sheets: list[SheetSummary] = []

    for worksheet in workbook.worksheets:
        row_iterator = worksheet.iter_rows(values_only=True)
        header_values = next(row_iterator, ())
        headers = [
            str(value).strip() if value is not None and str(value).strip() else f"column_{index + 1}"
            for index, value in enumerate(header_values)
        ]

        if not header_values:
            warnings.append(f"{worksheet.title} is empty")

        preview = []
        for row_index, row in enumerate(row_iterator):
            if row_index >= payload.max_preview_rows:
                break
            preview.append({headers[index]: row[index] for index in range(min(len(headers), len(row)))})

        sheets.append(
            SheetSummary(
                name=worksheet.title,
                rows=max(worksheet.max_row - 1, 0),
                columns=worksheet.max_column,
                headers=headers,
                preview=preview,
            )
        )

    workbook.close()

    return WorkbookInspectResponse(file_name=workbook_path.name, sheets=sheets, warnings=warnings)
