from pathlib import Path
from random import Random

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "public" / "demo-files"

WORKBOOKS = {
    "manufacturing": [
        ("Aarav_PnL_FY26_Q3.xlsx", "P&L", ["Revenue", "COGS", "Gross Profit", "Operating Expenses", "EBITDA", "Net Profit"]),
        ("Aarav_Balance_Sheet_FY26_Q3.xlsx", "Balance Sheet", ["Current Assets", "Inventory", "Cash", "Current Liabilities", "Total Debt", "Equity"]),
        ("Aarav_Cash_Flow_FY26_Q3.xlsx", "Cash Flow", ["Operating Cash Flow", "Investing Cash Flow", "Financing Cash Flow", "Net Cash Change"]),
        ("Aarav_Inventory_SKU_Q3.xlsx", "Inventory", ["Opening Qty", "Received Qty", "Issued Qty", "Book Qty", "Physical Qty", "Variance Value"]),
        ("Aarav_AP_Purchase_Ledger_Q3.xlsx", "AP Ledger", ["Invoice Amount", "Paid Amount", "Open Amount", "Age Days", "Duplicate Flag"]),
        ("Aarav_Sales_Register_Q3.xlsx", "Sales Register", ["Gross Sales", "Discount", "Net Sales", "GST", "Collection Days"]),
        ("Aarav_Budget_Vs_Actual_Q3.xlsx", "Budget", ["Budget", "Actual", "Variance", "Variance %"]),
    ],
    "retail": [
        ("Nila_Store_Sales_Q3.xlsx", "Sales Register", ["Gross Sales", "Discount", "Net Sales", "GST", "Policy Breach"]),
        ("Nila_GST_2B_3B_Q3.xlsx", "GST Summary", ["2B ITC", "3B Claimed", "Gap", "Supplier Count"]),
        ("Nila_Inventory_Ageing_Q3.xlsx", "Inventory", ["Stock Value", "0-30 Days", "31-90 Days", "90+ Days", "Dead Stock"]),
        ("Nila_PnL_Q3.xlsx", "P&L", ["Revenue", "COGS", "Gross Profit", "Operating Expenses", "EBITDA", "Net Profit"]),
        ("Nila_Balance_Sheet_Q3.xlsx", "Balance Sheet", ["Current Assets", "Inventory", "Cash", "Current Liabilities", "Total Debt", "Equity"]),
        ("Nila_Vendor_Ledger_Q3.xlsx", "AP Ledger", ["Invoice Amount", "Paid Amount", "Open Amount", "Age Days", "GSTIN Match"]),
        ("Nila_Budget_Vs_Actual_Q3.xlsx", "Budget", ["Budget", "Actual", "Variance", "Variance %"]),
    ],
    "saas": [
        ("VectorCloud_Subscription_CRM_Q3.xlsx", "Sales Register", ["MRR", "ARR", "Expansion", "Contraction", "Renewal Risk"]),
        ("VectorCloud_AR_Ageing_Q3.xlsx", "AR Ledger", ["Invoice Amount", "Collected", "Open Amount", "Age Days", "Risk Bucket"]),
        ("VectorCloud_Payroll_Summary_Q3.xlsx", "Payroll", ["Headcount", "Gross Pay", "Benefits", "Taxes", "Loaded Cost"]),
        ("VectorCloud_PnL_Q3.xlsx", "P&L", ["Revenue", "COGS", "Gross Profit", "Operating Expenses", "EBITDA", "Net Profit"]),
        ("VectorCloud_Balance_Sheet_Q3.xlsx", "Balance Sheet", ["Current Assets", "Cash", "Accounts Receivable", "Current Liabilities", "Debt", "Equity"]),
        ("VectorCloud_Cash_Flow_Q3.xlsx", "Cash Flow", ["Operating Cash Flow", "Investing Cash Flow", "Financing Cash Flow", "Net Cash Change"]),
        ("VectorCloud_Budget_Vs_Actual_Q3.xlsx", "Budget", ["Budget", "Actual", "Variance", "Variance %"]),
    ],
}

BASE_VALUES = {
    "manufacturing": 48_200_000,
    "retail": 28_600_000,
    "saas": 16_400_000,
}


def style_sheet(ws):
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows(min_row=1, max_row=4):
        for cell in row:
            cell.alignment = Alignment(vertical="center")
    for cell in ws[4]:
        cell.fill = PatternFill("solid", fgColor="123525")
        cell.font = Font(color="FFFFFF", bold=True)
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[chr(64 + col)].width = 18


def build_workbook(industry, file_name, doc_type, measures):
    rng = Random(f"{industry}-{file_name}")
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "FinLens Demo Input"
    summary["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor="0B1117")
    summary["A2"] = "Industry"
    summary["B2"] = industry.title()
    summary["A3"] = "Document Type"
    summary["B3"] = doc_type
    summary["A4"] = "Privacy Note"
    summary["B4"] = "Synthetic demo data for deterministic KPI, leakage, forecast, and lineage testing."
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 92

    data = wb.create_sheet("Source_Data")
    headers = ["Period", "Entity", "Cost Center", "Reference"] + measures
    data.append(["FinLens Demo Source Rows"])
    data.append([f"{doc_type} source data used by the demo workspace"])
    data.append([])
    data.append(headers)

    base = BASE_VALUES[industry]
    periods = ["Oct-2025", "Nov-2025", "Dec-2025", "Q3 FY26"]
    for idx in range(1, 31):
        row = [
            periods[idx % len(periods)],
            f"{industry.title()} Unit {1 + idx % 5}",
            f"CC-{100 + idx % 9}",
            f"REF-{industry[:3].upper()}-{idx:04d}",
        ]
        for measure_index, measure in enumerate(measures):
            if "%" in measure:
                row.append(round(rng.uniform(-0.18, 0.18), 4))
            elif "Flag" in measure or "Breach" in measure or "Match" in measure or "Risk" in measure:
                row.append(rng.choice(["Clear", "Watch", "Exception"]))
            elif "Days" in measure or "Headcount" in measure or "Count" in measure:
                row.append(rng.randint(8, 118))
            else:
                row.append(round(base * rng.uniform(0.006, 0.038), 0))
        data.append(row)

    table_ref = f"A4:{chr(64 + len(headers))}{data.max_row}"
    table = Table(displayName=f"{industry.title().replace(' ', '')}{doc_type.replace(' ', '')}Table", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True, showFirstColumn=False)
    data.add_table(table)
    style_sheet(data)

    lineage = wb.create_sheet("Lineage")
    lineage.append(["Metric", "Source Sheet", "Source Rows", "Formula / Rule", "Compute Engine"])
    lineage_rows = [
        ["Revenue / Amount", "Source_Data", "5:34", "SUM normalized amount fields by period", "finlens-demo-engine@0.2.0"],
        ["Health Score", "Source_Data", "5:34", "Completeness + mapping confidence + validation warnings", "data-quality@0.2.0"],
        ["Leak Finding", "Source_Data", "5:34", "Pattern checks for duplicate, abnormal, aged, or missing values", "leak-engine@0.2.0"],
    ]
    for row in lineage_rows:
        lineage.append(row)
    style_sheet(lineage)
    lineage.column_dimensions["D"].width = 54

    chart = LineChart()
    chart.title = f"{doc_type} Demo Trend"
    chart.y_axis.title = "Value"
    chart.x_axis.title = "Rows"
    value_col = 5
    values = Reference(data, min_col=value_col, min_row=4, max_row=16)
    chart.add_data(values, titles_from_data=True)
    summary.add_chart(chart, "A7")

    return wb


def main():
    for industry, files in WORKBOOKS.items():
        output_dir = OUT / industry
        output_dir.mkdir(parents=True, exist_ok=True)
        for file_name, doc_type, measures in files:
            workbook = build_workbook(industry, file_name, doc_type, measures)
            workbook.save(output_dir / file_name)


if __name__ == "__main__":
    main()
